"""Export the committed unclosed A/B notice snapshot as one XLSX worksheet."""

from __future__ import annotations

import json
import math
import os
import sqlite3
import unicodedata
from collections.abc import Mapping
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from mmis_connector.auth import MMISClientError


ROOT = Path(__file__).resolve().parents[2]
ENV_NAME = "UNCLOSED_AB_FAULT_NOTICES_XLSX_DIR"
SHEET_NAME = "AB級未處理故障通報"
FILE_PREFIX = "AB級未處理故障通報_程式產生"
FONT_NAME = "微軟正黑體"
FONT_SIZE = 12
LONG_TEXT_COLUMNS = {"故障現象", "備註"}
PENDING_WORK_ORDERS_FIELD = "待簽核工單"


class XlsxExportError(MMISClientError):
    """The SQLite sync committed, but XLSX export did not finish."""


def configured_output_directory() -> Path:
    load_dotenv(ROOT / ".env", override=False)
    value = os.getenv(ENV_NAME, "").strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
        value = value[1:-1].strip()
    if not value:
        raise XlsxExportError(f"未設定 {ENV_NAME}；請在專案 .env 設定 XLSX 輸出資料夾")
    directory = Path(value).expanduser()
    if not directory.is_absolute():
        raise XlsxExportError(f"{ENV_NAME} 必須是絕對路徑")
    if not directory.is_dir():
        raise XlsxExportError(f"{ENV_NAME} 指定的資料夾不存在")
    return directory


def read_snapshot(
    database_path: Path, table_name: str, expected_columns: tuple[str, ...],
) -> list[tuple[Any, ...]]:
    try:
        connection = sqlite3.connect(database_path.resolve().as_uri() + "?mode=ro", uri=True)
        try:
            actual = tuple(row[1] for row in connection.execute(f"PRAGMA table_info({table_name})"))
            if actual != expected_columns:
                raise XlsxExportError("本次 SQLite 資料表欄位與 XLSX 匯出規格不符")
            quoted = ", ".join(f'"{name}"' for name in expected_columns)
            return connection.execute(
                f"SELECT {quoted} FROM {table_name} ORDER BY rowid"
            ).fetchall()
        finally:
            connection.close()
    except sqlite3.Error as exc:
        raise XlsxExportError(f"無法讀取本次同步的 SQLite: {type(exc).__name__}") from exc


def _display_width(value: object) -> int:
    text = "TRUE" if value is True else "FALSE" if value is False else str(value) if value is not None else ""
    return max((sum(
        0 if unicodedata.combining(char) else
        2 if unicodedata.east_asian_width(char) in "FW" else 1
        for char in line
    ) for line in text.splitlines()), default=0)


def _line_count(value: object, usable_width: int) -> int:
    if value is None:
        return 1
    return sum(
        max(1, math.ceil(_display_width(line) / usable_width))
        for line in str(value).split("\n")
    )


def _display_pending_work_orders(value: object) -> str | None:
    if value is None:
        return None
    try:
        orders = json.loads(value)
    except (TypeError, ValueError) as exc:
        raise XlsxExportError("SQLite 待簽核工單不是有效的 JSON 陣列") from exc
    if not isinstance(orders, list) or any(not isinstance(order, str) for order in orders):
        raise XlsxExportError("SQLite 待簽核工單不是字串陣列")
    return ",".join(orders)


def build_workbook(
    rows: list[tuple[Any, ...]], columns: tuple[str, ...],
    boolean_columns: frozenset[str],
    *, row_height: float | None = None, vertical_alignment: str = "center",
    column_widths: Mapping[str, float] | None = None,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(columns)
    widths = [min(65 if name in LONG_TEXT_COLUMNS else 60,
                  max(10, _display_width(name) + 3)) for name in columns]
    header_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    data_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=False)
    for cell in sheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical=vertical_alignment)
    sheet.row_dimensions[1].height = row_height if row_height is not None else 24

    for row_number, row in enumerate(rows, start=2):
        if len(row) != len(columns):
            raise XlsxExportError("SQLite 資料列欄位數與 XLSX 匯出規格不符")
        sheet.append([None] * len(columns))
        for index, (name, value) in enumerate(zip(columns, row, strict=True), start=1):
            if name in boolean_columns and value is not None:
                if type(value) is not int or value not in (0, 1):
                    raise XlsxExportError(f"SQLite 布林欄位 {name} 含有無效值")
                value = bool(value)
                if name == "ATP故障":
                    value = "Y" if value else "N"
            cell = sheet.cell(row_number, index, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Preserve =, +, -, @ as literal text.
                cell.number_format = "@"
            cell.font = data_font
            cell.alignment = Alignment(horizontal="left", vertical=vertical_alignment)
            if value is not None:
                widths[index - 1] = min(
                    65 if name in LONG_TEXT_COLUMNS else 60,
                    max(widths[index - 1], _display_width(value) + 3),
                )

    for index, (name, width) in enumerate(zip(columns, widths, strict=True), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            column_widths[name] if column_widths is not None and name in column_widths else width
        )
    note_index = columns.index("備註") + 1 if "備註" in columns else None
    for row_number in range(2, sheet.max_row + 1):
        height_lines = 1
        for index, name in enumerate(columns, start=1):
            cell = sheet.cell(row_number, index)
            usable = max(1, math.floor(sheet.column_dimensions[get_column_letter(index)].width - 3))
            lines = _line_count(cell.value, usable)
            if name in LONG_TEXT_COLUMNS or lines > 1:
                cell.alignment = Alignment(
                    horizontal="left", vertical=vertical_alignment, wrap_text=True
                )
            height_lines = max(height_lines, lines)
        if row_height is not None:
            note = sheet.cell(row_number, note_index).value if note_index is not None else None
            note_lines = str(note).count("\n") + 1 if note else 1
            sheet.row_dimensions[row_number].height = min(409.5, row_height * note_lines)
        else:
            sheet.row_dimensions[row_number].height = min(409.5, 18 * height_lines + 6)
    return workbook


def export_xlsx(
    database_path: Path, table_name: str, columns: tuple[str, ...],
    boolean_columns: frozenset[str], *, output_directory: Path | None = None,
    now: datetime | None = None, expected_records: int | None = None,
    excluded_columns: frozenset[str] = frozenset(),
    row_height: float | None = None, vertical_alignment: str = "center",
    column_widths: Mapping[str, float] | None = None,
    sort_columns: tuple[str, ...] = (),
) -> dict[str, Any]:
    directory = output_directory if output_directory is not None else configured_output_directory()
    if not directory.is_dir():
        raise XlsxExportError("XLSX 輸出資料夾不存在")
    rows = read_snapshot(database_path, table_name, columns)
    for row in rows:
        for index, name in enumerate(columns):
            value = row[index]
            if name in boolean_columns and value is not None and (
                type(value) is not int or value not in (0, 1)
            ):
                raise XlsxExportError(f"SQLite 布林欄位 {name} 含有無效值")
    if expected_records is not None and len(rows) != expected_records:
        raise XlsxExportError("本次 SQLite 筆數與 XLSX 匯出筆數不符")
    if not excluded_columns.issubset(columns):
        raise XlsxExportError("XLSX 排除欄位不在 SQLite schema 中")
    if not set(sort_columns).issubset(columns):
        raise XlsxExportError("XLSX 排序欄位不在 SQLite schema 中")
    visible_indices = [index for index, name in enumerate(columns) if name not in excluded_columns]
    visible_columns = tuple(columns[index] for index in visible_indices)
    if column_widths is not None and not set(column_widths).issubset(visible_columns):
        raise XlsxExportError("XLSX 欄寬設定欄位不在匯出欄位中")
    sort_indices = tuple(columns.index(name) for name in sort_columns)
    if sort_indices:
        rows.sort(key=lambda row: tuple(
            (row[index] is None, str(row[index]) if row[index] is not None else "")
            for index in sort_indices
        ))
    visible_rows = [
        tuple(
            _display_pending_work_orders(row[index])
            if columns[index] == PENDING_WORK_ORDERS_FIELD else row[index]
            for index in visible_indices
        )
        for row in rows
    ]
    workbook = build_workbook(
        visible_rows, visible_columns, boolean_columns,
        row_height=row_height, vertical_alignment=vertical_alignment,
        column_widths=column_widths,
    )
    return save_workbook(
        workbook, directory, file_prefix=FILE_PREFIX,
        records_exported=len(rows), columns_exported=len(visible_columns), now=now,
    )


def save_workbook(
    workbook: Workbook, directory: Path, *, file_prefix: str,
    records_exported: int, columns_exported: int,
    now: datetime | None = None,
) -> dict[str, Any]:
    date = (now or datetime.now(ZoneInfo("Asia/Taipei"))).astimezone(
        ZoneInfo("Asia/Taipei")
    ).strftime("%m%d")
    base = f"{file_prefix}{date}"
    output_file: Path | None = None
    export_done = False
    try:
        suffix = 1
        while True:
            filename = f"{base}{'_' + str(suffix) if suffix > 1 else ''}.xlsx"
            candidate = directory / filename
            try:
                handle = candidate.open("xb")
            except FileExistsError:
                suffix += 1
                continue
            output_file = candidate
            break
        with handle:
            workbook.save(handle)
        export_done = True
        return {
            "status": "success", "records_exported": records_exported,
            "columns_exported": columns_exported, "output_file": str(output_file),
        }
    except OSError as exc:
        raise XlsxExportError(f"XLSX 寫入失敗: {type(exc).__name__}") from exc
    finally:
        workbook.close()
        if output_file is not None and not export_done:
            output_file.unlink(missing_ok=True)
