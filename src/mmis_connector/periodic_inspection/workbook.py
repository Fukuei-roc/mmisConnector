from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Protocol

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from ..auth import MMISClientError


SHEET_NAME = "最新檢修工單"
HEADERS = ("車組/車號", "檢修級別", "工作單", "檢修日期")
DEFAULT_WORKBOOK = Path(r"C:\Users\NMMIS\Downloads\軔管日期統計\配屬車號總表.xlsx")
MAX_ATTEMPTS = 3


class PeriodicInspectionQuery(Protocol):
    def run(self, vehicle: str) -> dict[str, str]: ...


@dataclass(frozen=True)
class BatchResult:
    processed: int
    skipped: int
    failed_rows: tuple[int, ...]

    def as_dict(self) -> dict[str, object]:
        return {
            "success": not self.failed_rows,
            "processed": self.processed,
            "skipped": self.skipped,
            "failed": len(self.failed_rows),
            "failed_rows": list(self.failed_rows),
        }


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def load_validated_workbook(path: Path) -> Workbook:
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise MMISClientError("找不到可讀取的 .xlsx 活頁簿")
    try:
        workbook = load_workbook(path)
    except (OSError, ValueError, RuntimeError) as exc:
        raise MMISClientError("無法讀取 Excel 活頁簿") from exc
    except Exception as exc:
        raise MMISClientError("Excel 活頁簿格式無效") from exc
    if SHEET_NAME not in workbook.sheetnames:
        raise MMISClientError(f"Excel 缺少工作表：{SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    if tuple(sheet.cell(1, column).value for column in range(1, 5)) != HEADERS:
        raise MMISClientError("Excel A 至 D 欄標題不符")
    for row_number in range(2, sheet.max_row + 1):
        vehicle = sheet.cell(row_number, 1)
        if vehicle.data_type == "f" or isinstance(vehicle.value, (bool, complex)):
            raise MMISClientError(f"Excel 第 {row_number} 列車組/車號格式無效")
    return workbook


def _save_atomic(workbook: Workbook, path: Path) -> None:
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}.", suffix=".xlsx", dir=path.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    except OSError as exc:
        raise MMISClientError("無法儲存 Excel 活頁簿；已停止處理") from exc
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def fill_latest_periodic_inspection_work_orders(
    workbook: Workbook, path: Path, query_factory: Callable[[], PeriodicInspectionQuery]
) -> BatchResult:
    sheet = workbook[SHEET_NAME]
    processed = 0
    skipped = 0
    failed_rows: list[int] = []
    query: PeriodicInspectionQuery | None = None
    for row_number in range(2, sheet.max_row + 1):
        vehicle = sheet.cell(row_number, 1).value
        output_cells = [sheet.cell(row_number, column) for column in range(2, 5)]
        if not _has_value(vehicle) or all(_has_value(cell.value) for cell in output_cells):
            skipped += 1
            continue
        values: tuple[str, ...] | None = None
        for _ in range(MAX_ATTEMPTS):
            try:
                if query is None:
                    query = query_factory()
                result = query.run(str(vehicle).strip())
            except MMISClientError:
                query = None
                continue
            values = tuple(str(result[field]).strip() for field in HEADERS[1:])
            if not all(values):
                raise ValueError("查詢結果缺少必要欄位")
            break
        if values is None:
            failed_rows.append(row_number)
            continue
        for cell, value in zip(output_cells, values):
            cell.value = value
            cell.data_type = "s"
        _save_atomic(workbook, path)
        processed += 1
    return BatchResult(processed, skipped, tuple(failed_rows))
