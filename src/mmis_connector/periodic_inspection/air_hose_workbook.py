from __future__ import annotations

import hashlib
import json
import os
import tempfile
from collections.abc import Callable, Mapping
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
import requests

from ..auth import MMISClientError
from ..daily_inspection.reader import normalize_work_order


SOURCE_SHEET = "最新檢修工單"
SOURCE_HEADERS = ("車組/車號", "檢修級別", "工作單", "檢修日期")
OUTPUT_SHEET = "空氣軟管日期"
OUTPUT_HEADERS = (
    "車組/車號", "車型", "檢修級別", "工作單", "檢修日期", "檢查項目", "量測值",
)
PROGRESS_SHEET = "_air_hose_progress"
PROGRESS_HEADERS = (
    "source_row", "work_order", "source_digest", "status", "vehicle_model",
    "output_start", "record_count", "output_digest",
)
SCHEMA_VERSION = 1
MAX_ATTEMPTS = 3
DEFAULT_SOURCE = Path(r"C:\Users\NMMIS\Downloads\軔管日期統計\配屬車號總表.xlsx")
DEFAULT_OUTPUT = Path(r"C:\Users\NMMIS\Downloads\軔管日期統計\空氣軟管日期.xlsx")


class AirHoseReader(Protocol):
    def run_air_hose_dates(self, work_order: str) -> dict[str, Any]: ...


@dataclass(frozen=True)
class SourceOrder:
    row_number: int
    values: tuple[object, object, object, object]
    work_order: str
    digest: str


@dataclass(frozen=True)
class CompletedOrder:
    source: SourceOrder
    status: str
    vehicle_model: str | None
    rows: tuple[tuple[object, ...], ...]


@dataclass(frozen=True)
class BatchResult:
    processed: int
    skipped: int
    unsupported_rows: tuple[int, ...]
    empty_rows: tuple[int, ...]
    failed_rows: tuple[int, ...]
    failed_work_orders: tuple[str, ...]
    failed_error_types: tuple[str, ...]
    measurement_rows: int

    @property
    def unsupported(self) -> int:
        return len(self.unsupported_rows)

    @property
    def empty(self) -> int:
        return len(self.empty_rows)

    def as_dict(self) -> dict[str, object]:
        return {
            "success": not self.failed_rows,
            "processed": self.processed,
            "skipped": self.skipped,
            "unsupported": self.unsupported,
            "unsupported_rows": list(self.unsupported_rows),
            "empty": self.empty,
            "empty_rows": list(self.empty_rows),
            "measurements": self.measurement_rows,
            "failed": len(self.failed_rows),
            "failed_rows": list(self.failed_rows),
            "failed_work_orders": list(self.failed_work_orders),
            "failed_error_types": list(self.failed_error_types),
        }


def _canonical(value: object) -> tuple[str, str | None]:
    if value is None:
        return ("empty", None)
    if isinstance(value, datetime):
        if value.time() == time.min:
            return ("date", value.date().isoformat())
        return ("datetime", value.isoformat())
    if isinstance(value, date):
        return ("date", value.isoformat())
    if isinstance(value, time):
        return ("time", value.isoformat())
    if isinstance(value, str):
        return ("text", value)
    if type(value) in (int, float, Decimal):
        return ("number", str(Decimal(str(value)).normalize()))
    raise MMISClientError("Excel 包含不支援的儲存格資料型態")


def _hash(value: object) -> str:
    encoded = json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _row_digest(row_number: int, values: tuple[object, ...]) -> str:
    return _hash([row_number, [_canonical(value) for value in values]])


def _output_digest(rows: tuple[tuple[object, ...], ...]) -> str:
    return _hash([
        [_canonical("" if index == 6 and value is None else value)
         for index, value in enumerate(row)]
        for row in rows
    ])


def _present(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _load_source(path: Path) -> tuple[tuple[SourceOrder, ...], str]:
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise MMISClientError("找不到來源 Excel 活頁簿")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise MMISClientError("無法讀取來源 Excel 活頁簿") from exc
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise MMISClientError(f"來源 Excel 缺少工作表：{SOURCE_SHEET}")
        sheet = workbook[SOURCE_SHEET]
        if tuple(sheet.cell(1, column).value for column in range(1, 5)) != SOURCE_HEADERS:
            raise MMISClientError("來源 Excel A 至 D 欄標題不符")
        orders: list[SourceOrder] = []
        for row_number in range(2, sheet.max_row + 1):
            cells = tuple(sheet.cell(row_number, column) for column in range(1, 5))
            values = tuple(cell.value for cell in cells)
            if not any(_present(value) for value in values):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise MMISClientError(f"來源 Excel 第 {row_number} 列含公式")
            if any(isinstance(value, (bool, complex)) for value in values):
                raise MMISClientError(f"來源 Excel 第 {row_number} 列資料型態無效")
            if not _present(values[2]):
                raise MMISClientError(f"來源 Excel 第 {row_number} 列缺少工作單")
            work_order = normalize_work_order(str(values[2]))
            typed_values = (values[0], values[1], values[2], values[3])
            orders.append(SourceOrder(
                row_number, typed_values, work_order,
                _row_digest(row_number, typed_values),
            ))
        snapshot = _hash([
            SOURCE_SHEET, SOURCE_HEADERS,
            [(order.row_number, order.digest) for order in orders],
        ])
        return tuple(orders), snapshot
    finally:
        workbook.close()


def _load_output(path: Path) -> Workbook:
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise MMISClientError("找不到目標 Excel 活頁簿")
    try:
        workbook = load_workbook(path)
    except Exception as exc:
        raise MMISClientError("無法讀取目標 Excel 活頁簿") from exc
    if set(workbook.sheetnames) not in ({OUTPUT_SHEET}, {OUTPUT_SHEET, PROGRESS_SHEET}):
        workbook.close()
        raise MMISClientError("目標 Excel 工作表結構不符")
    sheet = workbook[OUTPUT_SHEET]
    if tuple(sheet.cell(1, column).value for column in range(1, 8)) != OUTPUT_HEADERS:
        workbook.close()
        raise MMISClientError("目標 Excel A 至 G 欄標題不符")
    if any(sheet.cell(1, column).value is not None for column in range(8, sheet.max_column + 1)):
        workbook.close()
        raise MMISClientError("目標 Excel 含有預期外欄位")
    return workbook


def _create_progress(workbook: Workbook, source_fingerprint: str) -> None:
    sheet = workbook.create_sheet(PROGRESS_SHEET)
    sheet.cell(1, 1, "schema_version")
    sheet.cell(1, 2, SCHEMA_VERSION)
    sheet.cell(2, 1, "source_fingerprint")
    sheet.cell(2, 2, source_fingerprint)
    for column, header in enumerate(PROGRESS_HEADERS, start=1):
        sheet.cell(4, column, header)
    sheet.sheet_state = "hidden"


def _load_completed(
    workbook: Workbook, source_orders: tuple[SourceOrder, ...], source_fingerprint: str,
) -> dict[int, CompletedOrder]:
    visible = workbook[OUTPUT_SHEET]
    if PROGRESS_SHEET not in workbook:
        if visible.max_row > 1 and any(
            visible.cell(row, column).value is not None
            for row in range(2, visible.max_row + 1)
            for column in range(1, visible.max_column + 1)
        ):
            raise MMISClientError("目標 Excel 已有資料但缺少續跑紀錄")
        _create_progress(workbook, source_fingerprint)
        return {}
    progress = workbook[PROGRESS_SHEET]
    if (
        progress.sheet_state != "hidden"
        or progress.cell(1, 1).value != "schema_version"
        or type(progress.cell(1, 2).value) is not int
        or progress.cell(1, 2).value != SCHEMA_VERSION
        or progress.cell(2, 1).value != "source_fingerprint"
        or progress.cell(2, 2).value != source_fingerprint
        or tuple(progress.cell(4, column).value for column in range(1, 9)) != PROGRESS_HEADERS
    ):
        raise MMISClientError("目標 Excel 續跑紀錄與來源或版本不相符")
    by_row = {order.row_number: order for order in source_orders}
    completed: dict[int, CompletedOrder] = {}
    next_visible_row = 2
    previous_source_row = 1
    for progress_row in range(5, progress.max_row + 1):
        source_row, number, source_digest, status, model, start, count, digest = (
            progress.cell(progress_row, column).value for column in range(1, 9)
        )
        source = by_row.get(source_row) if type(source_row) is int else None
        if (
            source is None or source_row <= previous_source_row
            or number != source.work_order or source_digest != source.digest
            or status not in ("complete", "empty", "unsupported")
            or type(count) is not int or count < 0
            or (status == "complete") != (count > 0)
            or (status == "unsupported") != (model is None)
            or (status != "unsupported" and (not isinstance(model, str) or not model))
            or start != (next_visible_row if count else None)
            or not isinstance(digest, str)
        ):
            raise MMISClientError("目標 Excel 續跑紀錄內容無效")
        rows = tuple(
            tuple(visible.cell(row, column).value for column in range(1, 8))
            for row in range(next_visible_row, next_visible_row + count)
        )
        if digest != _output_digest(rows) or any(
            row[:1] != (source.values[0],)
            or row[1] != model
            or row[2:5] != (source.values[1], source.values[2], source.values[3])
            for row in rows
        ):
            raise MMISClientError("目標 Excel 量測資料與續跑紀錄不一致")
        completed[source_row] = CompletedOrder(source, status, model, rows)
        previous_source_row = source_row
        next_visible_row += count
    if visible.max_row != next_visible_row - 1 or any(
        visible.cell(row, column).value is not None
        for row in range(2, visible.max_row + 1)
        for column in range(8, visible.max_column + 1)
    ):
        raise MMISClientError("目標 Excel 存在未記錄的量測資料")
    return completed


def _parse_result(source: SourceOrder, result: object) -> CompletedOrder:
    if not isinstance(result, Mapping):
        raise MMISClientError("工作單查詢結果格式無效")
    supported = result.get("supported")
    model = result.get("車型")
    records = result.get("records")
    count = result.get("count")
    if (
        result.get("success") is not True
        or result.get("工作單") != source.work_order
        or type(supported) is not bool
        or not isinstance(records, list)
        or type(count) is not int
        or count != len(records)
        or (supported and (not isinstance(model, str) or not model.strip()))
        or (not supported and (count != 0 or model is not None))
    ):
        raise MMISClientError("工作單查詢結果與來源資料不相符")
    rows: list[tuple[object, ...]] = []
    for record in records:
        if (
            not isinstance(record, Mapping)
            or not isinstance(record.get("檢查項目"), str)
            or not record["檢查項目"].strip()
            or not isinstance(record.get("量測值"), str)
        ):
            raise MMISClientError("工作單量測資料欄位無效")
        rows.append((
            source.values[0], model, source.values[1], source.values[2],
            source.values[3], record["檢查項目"], record["量測值"],
        ))
    status = "unsupported" if not supported else "complete" if rows else "empty"
    return CompletedOrder(source, status, model, tuple(rows))


def _render(workbook: Workbook, completed: dict[int, CompletedOrder]) -> None:
    visible = workbook[OUTPUT_SHEET]
    progress = workbook[PROGRESS_SHEET]
    if visible.max_row > 1:
        visible.delete_rows(2, visible.max_row - 1)
    if progress.max_row > 4:
        progress.delete_rows(5, progress.max_row - 4)
    for item in (completed[row] for row in sorted(completed)):
        start = visible.max_row + 1 if item.rows else None
        for values in item.rows:
            visible.append(list(values))
            output_row = visible.max_row
            for column, value in enumerate(values, start=1):
                if isinstance(value, str):
                    visible.cell(output_row, column).data_type = "s"
        progress.append((
            item.source.row_number, item.source.work_order, item.source.digest,
            item.status, item.vehicle_model, start, len(item.rows),
            _output_digest(item.rows),
        ))


def _save_atomic(workbook: Workbook, path: Path) -> None:
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}.", suffix=".xlsx", dir=path.parent, delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    except OSError as exc:
        raise MMISClientError("無法儲存目標 Excel 活頁簿；已保留先前結果") from exc
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


@contextmanager
def _output_lock(path: Path):
    lock_path = path.with_name(f".{path.name}.lock")
    try:
        handle = lock_path.open("a+b")
    except OSError as exc:
        raise MMISClientError("無法鎖定目標 Excel 活頁簿") from exc
    with handle:
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"\0")
            handle.flush()
        handle.seek(0)
        try:
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl

                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            raise MMISClientError("另一個空氣軟管日期查詢程序正在寫入目標檔案") from exc
        try:
            yield
        finally:
            handle.seek(0)
            if os.name == "nt":
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def query_air_hose_dates_to_workbook(
    source_path: Path, output_path: Path, reader_factory: Callable[[], AirHoseReader],
) -> BatchResult:
    source_orders, fingerprint = _load_source(source_path)
    if output_path.suffix.lower() != ".xlsx" or not output_path.is_file():
        raise MMISClientError("找不到目標 Excel 活頁簿")
    with _output_lock(output_path):
        workbook = _load_output(output_path)
        try:
            completed = _load_completed(workbook, source_orders, fingerprint)
            processed = skipped = 0
            failed_rows: list[int] = []
            failed_work_orders: list[str] = []
            failed_error_types: list[str] = []
            reader: AirHoseReader | None = None
            for source in source_orders:
                if source.row_number in completed:
                    skipped += 1
                    continue
                item: CompletedOrder | None = None
                last_error: MMISClientError | None = None
                for _ in range(MAX_ATTEMPTS):
                    try:
                        if reader is None:
                            reader = reader_factory()
                        item = _parse_result(
                            source, reader.run_air_hose_dates(source.work_order),
                        )
                    except MMISClientError as exc:
                        last_error = exc
                        reader = None
                        continue
                    break
                if item is None:
                    failed_rows.append(source.row_number)
                    failed_work_orders.append(source.work_order)
                    cause = last_error.__cause__ if last_error is not None else None
                    failed_error_types.append(
                        type(cause).__name__
                        if isinstance(cause, requests.RequestException)
                        else type(last_error).__name__
                        if last_error is not None
                        else "UnknownError"
                    )
                    continue
                completed[source.row_number] = item
                _render(workbook, completed)
                _save_atomic(workbook, output_path)
                processed += 1
            unsupported_rows = tuple(
                row for row, item in sorted(completed.items())
                if item.status == "unsupported"
            )
            empty_rows = tuple(
                row for row, item in sorted(completed.items())
                if item.status == "empty"
            )
            return BatchResult(
                processed, skipped, unsupported_rows, empty_rows,
                tuple(failed_rows), tuple(failed_work_orders),
                tuple(failed_error_types),
                sum(len(item.rows) for item in completed.values()),
            )
        finally:
            workbook.close()
