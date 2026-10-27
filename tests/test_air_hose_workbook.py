from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest
import requests
from openpyxl import Workbook, load_workbook

from mmis_connector.auth import MMISClientError
from mmis_connector.periodic_inspection import air_hose_workbook as batch
from tools.mmis_development import auto_query_air_hose_dates as cli


def make_files(
    tmp_path: Path, rows: list[tuple[object, object, object, object]],
) -> tuple[Path, Path]:
    source_path = tmp_path / "source.xlsx"
    source = Workbook()
    source_sheet = source.active
    source_sheet.title = batch.SOURCE_SHEET
    source_sheet.append(batch.SOURCE_HEADERS)
    for row in rows:
        source_sheet.append(row)
    source.save(source_path)

    output_path = tmp_path / "output.xlsx"
    output = Workbook()
    output_sheet = output.active
    output_sheet.title = batch.OUTPUT_SHEET
    output_sheet.append(batch.OUTPUT_HEADERS)
    output.save(output_path)
    return source_path, output_path


def query_result(
    vehicle: str, work_order: str, count: int = 1, *, supported: bool = True,
    measurement: str = "112/7",
) -> dict[str, object]:
    return {
        "success": True,
        "supported": supported,
        "工作單": work_order,
        "車組/車號": vehicle,
        "車型": "800型" if supported else None,
        "count": count,
        "records": [
            {"檢查項目": f"空氣軟管 {index + 1}", "量測值": measurement}
            for index in range(count)
        ],
    }


def visible_rows(path: Path) -> list[tuple[object, ...]]:
    workbook = load_workbook(path)
    try:
        return list(workbook[batch.OUTPUT_SHEET].values)
    finally:
        workbook.close()


def test_expands_all_measurements_and_skips_completed_orders_on_rerun(tmp_path: Path) -> None:
    first, second = "115-3A-100-001", "115-3A-101-001"
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", first, "2026/10/20"),
        ("EMU802", "2A", second, datetime(2026, 10, 21)),
    ])
    calls: list[str] = []

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            calls.append(number)
            if number == first:
                return query_result("EMU801", number, 12)
            return query_result("EMU802", number, 13, measurement="")

    result = batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert result.as_dict() == {
        "success": True, "processed": 2, "skipped": 0,
        "unsupported": 0, "unsupported_rows": [],
        "empty": 0, "empty_rows": [], "measurements": 25,
        "failed": 0, "failed_rows": [], "failed_work_orders": [],
        "failed_error_types": [],
    }
    assert calls == [first, second]
    assert [
        row[5:] for row in visible_rows(output)[1:13]
    ] == [
        (f"空氣軟管 {index}", "112/7") for index in range(1, 13)
    ]
    assert [
        row[5:] for row in visible_rows(output)[13:]
    ] == [
        (f"空氣軟管 {index}", None) for index in range(1, 14)
    ]
    workbook = load_workbook(output)
    try:
        sheet = workbook[batch.OUTPUT_SHEET]
        assert sheet.max_row == 26
        assert tuple(sheet[2][column].value for column in range(7)) == (
            "EMU801", "800型", "3A", first, "2026/10/20", "空氣軟管 1", "112/7",
        )
        assert tuple(sheet[14][column].value for column in range(7)) == (
            "EMU802", "800型", "2A", second, datetime(2026, 10, 21),
            "空氣軟管 1", None,
        )
        assert workbook[batch.PROGRESS_SHEET].sheet_state == "hidden"
    finally:
        workbook.close()

    def fail_if_called():
        raise AssertionError("Completed work orders were queried again")

    resumed = batch.query_air_hose_dates_to_workbook(source, output, fail_if_called)
    assert resumed.processed == 0
    assert resumed.skipped == 2
    assert resumed.measurement_rows == 25


def test_zero_and_unsupported_results_are_completed_without_visible_rows(tmp_path: Path) -> None:
    first, second = "115-3A-100-001", "115-3A-101-001"
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", first, "2026/10/20"),
        ("EMU901", "2A", second, "2026/10/21"),
    ])

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            if number == first:
                return query_result("EMU801", number, 0)
            return query_result("EMU901", number, 0, supported=False)

    result = batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert result.processed == 2
    assert result.empty == 1
    assert result.unsupported == 1
    assert visible_rows(output) == [batch.OUTPUT_HEADERS]
    assert batch.query_air_hose_dates_to_workbook(
        source, output, lambda: pytest.fail("No query expected"),
    ).skipped == 2


def test_retries_with_a_fresh_reader(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    created = 0

    class Reader:
        def __init__(self, should_fail: bool) -> None:
            self.should_fail = should_fail

        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            if self.should_fail:
                raise MMISClientError("MMIS 網路請求失敗")
            return query_result("EMU801", number)

    def factory() -> Reader:
        nonlocal created
        created += 1
        return Reader(created == 1)

    result = batch.query_air_hose_dates_to_workbook(source, output, factory)
    assert result.processed == 1
    assert created == 2
    assert len(visible_rows(output)) == 2


def test_exhausted_failure_continues_then_resumes_in_source_order(tmp_path: Path) -> None:
    first, second = "115-3A-100-001", "115-3A-101-001"
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", first, "2026/10/20"),
        ("EMU802", "3A", second, "2026/10/21"),
    ])
    attempts: list[str] = []

    class InitialReader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            attempts.append(number)
            if number == first:
                raise MMISClientError("temporary outage") from requests.ConnectionError(
                    "sensitive transport detail"
                )
            return query_result("EMU802", number)

    initial = batch.query_air_hose_dates_to_workbook(source, output, InitialReader)
    assert initial.failed_rows == (2,)
    assert initial.failed_work_orders == (first,)
    assert initial.failed_error_types == ("ConnectionError",)
    assert initial.processed == 1
    assert attempts == [first, first, first, second]
    assert "sensitive transport detail" not in str(initial.as_dict())
    assert visible_rows(output)[1][3] == second

    resumed_calls: list[str] = []

    class ResumedReader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            resumed_calls.append(number)
            return query_result("EMU801", number)

    resumed = batch.query_air_hose_dates_to_workbook(source, output, ResumedReader)
    assert resumed_calls == [first]
    assert resumed.processed == 1
    assert resumed.skipped == 1
    assert [row[3] for row in visible_rows(output)[1:]] == [first, second]


def test_resume_summary_keeps_prior_unsupported_and_empty_counts(tmp_path: Path) -> None:
    unsupported_order, empty_order = "115-3A-100-001", "115-3A-101-001"
    source, output = make_files(tmp_path, [
        ("EMU901", "3A", unsupported_order, "2026/10/20"),
        ("EMU801", "3A", empty_order, "2026/10/21"),
    ])

    class FirstRunReader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            if number == unsupported_order:
                return query_result("EMU901", number, count=0, supported=False)
            raise MMISClientError("network failure") from requests.ConnectionError()

    first = batch.query_air_hose_dates_to_workbook(source, output, FirstRunReader)
    assert first.unsupported == 1
    assert first.unsupported_rows == (2,)
    assert first.empty == 0
    assert first.failed_work_orders == (empty_order,)

    class ResumedReader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            assert number == empty_order
            return query_result("EMU801", number, count=0)

    resumed = batch.query_air_hose_dates_to_workbook(source, output, ResumedReader)
    assert resumed.skipped == 1
    assert resumed.processed == 1
    assert resumed.unsupported == 1
    assert resumed.unsupported_rows == (2,)
    assert resumed.empty == 1
    assert resumed.empty_rows == (3,)


def test_processes_duplicate_work_orders_as_distinct_source_rows(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
        ("EMU802", "3A", "115-3A-100-001", "2026/10/21"),
    ])
    calls: list[str] = []

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            calls.append(number)
            return query_result("EMU801", number)

    result = batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert result.processed == 2
    assert calls == ["115-3A-100-001", "115-3A-100-001"]
    assert [row[4] for row in visible_rows(output)[1:]] == [
        "2026/10/20", "2026/10/21",
    ]
    assert [row[0] for row in visible_rows(output)[1:]] == ["EMU801", "EMU802"]
    resumed = batch.query_air_hose_dates_to_workbook(
        source, output, lambda: pytest.fail("Completed source rows were queried"),
    )
    assert resumed.skipped == 2


def test_rejects_changed_source_and_tampered_visible_values(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            return query_result("EMU801", number)

    batch.query_air_hose_dates_to_workbook(source, output, Reader)
    source_book = load_workbook(source)
    source_book[batch.SOURCE_SHEET]["D2"] = "2026/10/22"
    source_book.save(source)
    with pytest.raises(MMISClientError, match="來源|續跑"):
        batch.query_air_hose_dates_to_workbook(source, output, Reader)
    source_book[batch.SOURCE_SHEET]["D2"] = "2026/10/20"
    source_book.save(source)
    source_book.close()

    output_book = load_workbook(output)
    output_book[batch.OUTPUT_SHEET]["G2"] = "altered"
    output_book.save(output)
    output_book.close()
    with pytest.raises(MMISClientError, match="不一致"):
        batch.query_air_hose_dates_to_workbook(source, output, Reader)


def test_refuses_output_data_without_checkpoint(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    book = load_workbook(output)
    book[batch.OUTPUT_SHEET].append((
        "EMU801", "800型", "3A", "115-3A-100-001", "2026/10/20",
        "空氣軟管", "112/7",
    ))
    book.save(output)
    book.close()
    with pytest.raises(MMISClientError, match="缺少續跑紀錄"):
        batch.query_air_hose_dates_to_workbook(
            source, output, lambda: pytest.fail("No query expected"),
        )


def test_replace_failure_keeps_last_saved_workbook(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    before = output.read_bytes()

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            return query_result("EMU801", number)

    def fail_replace(*_args: object) -> None:
        raise PermissionError("Excel open")

    monkeypatch.setattr(batch.os, "replace", fail_replace)
    with pytest.raises(MMISClientError, match="儲存"):
        batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert output.read_bytes() == before
    assert list(tmp_path.glob(".output.*.xlsx")) == []


def test_replace_failure_after_a_success_keeps_progress_and_resumes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    first, second = "115-3A-100-001", "115-3A-101-001"
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", first, "2026/10/20"),
        ("EMU802", "3A", second, "2026/10/21"),
    ])
    original_replace = batch.os.replace
    saves = 0

    def fail_second_replace(src: object, dst: object) -> None:
        nonlocal saves
        saves += 1
        if saves == 2:
            raise PermissionError("Excel open")
        original_replace(src, dst)

    monkeypatch.setattr(batch.os, "replace", fail_second_replace)

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            return query_result("EMU801" if number == first else "EMU802", number)

    with pytest.raises(MMISClientError, match="儲存"):
        batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert len(visible_rows(output)) == 2

    monkeypatch.setattr(batch.os, "replace", original_replace)
    resumed_calls: list[str] = []

    class ResumedReader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            resumed_calls.append(number)
            return query_result("EMU802", number)

    result = batch.query_air_hose_dates_to_workbook(source, output, ResumedReader)
    assert resumed_calls == [second]
    assert result.skipped == 1
    assert result.processed == 1
    assert [row[3] for row in visible_rows(output)[1:]] == [first, second]


def test_allows_blank_optional_source_fields(tmp_path: Path) -> None:
    work_order = "115-3A-100-001"
    source, output = make_files(tmp_path, [
        (None, None, work_order, None),
    ])

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            return query_result("EMU801", number)

    result = batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert result.processed == 1
    assert visible_rows(output)[1] == (
        None, "800型", None, work_order, None, "空氣軟管 1", "112/7",
    )


@pytest.mark.parametrize("target", ["source", "output"])
def test_rejects_invalid_headers_before_query(tmp_path: Path, target: str) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    path = source if target == "source" else output
    book = load_workbook(path)
    sheet = book[batch.SOURCE_SHEET if target == "source" else batch.OUTPUT_SHEET]
    sheet.cell(1, 1, "錯誤表頭")
    book.save(path)
    book.close()

    original_output = output.read_bytes()
    with pytest.raises(MMISClientError, match="標題"):
        batch.query_air_hose_dates_to_workbook(
            source, output, lambda: pytest.fail("No query expected"),
        )
    assert output.read_bytes() == original_output


def test_rejects_formula_in_source_before_query(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    book = load_workbook(source)
    book[batch.SOURCE_SHEET]["D2"] = "=TODAY()"
    book.save(source)
    book.close()
    original_output = output.read_bytes()

    with pytest.raises(MMISClientError, match="公式"):
        batch.query_air_hose_dates_to_workbook(
            source, output, lambda: pytest.fail("No query expected"),
        )
    assert output.read_bytes() == original_output


def test_malformed_query_result_remains_retryable(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    calls = 0

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            nonlocal calls
            calls += 1
            result = query_result("EMU801", number)
            result["count"] = 2
            return result

    result = batch.query_air_hose_dates_to_workbook(source, output, Reader)
    assert calls == 3
    assert result.failed_rows == (2,)
    assert visible_rows(output) == [batch.OUTPUT_HEADERS]


def test_second_writer_cannot_enter_the_same_output(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])
    with batch._output_lock(output):
        with pytest.raises(MMISClientError, match="另一個"):
            batch.query_air_hose_dates_to_workbook(
                source, output, lambda: pytest.fail("No query expected"),
            )


def test_query_strings_remain_literal_excel_text(tmp_path: Path) -> None:
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", "115-3A-100-001", "2026/10/20"),
    ])

    class Reader:
        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            return query_result("EMU801", number, measurement="=2+3")

    batch.query_air_hose_dates_to_workbook(source, output, Reader)
    book = load_workbook(output)
    try:
        cell = book[batch.OUTPUT_SHEET]["G2"]
        assert cell.value == "=2+3"
        assert cell.data_type == "s"
    finally:
        book.close()


def test_cli_returns_nonzero_with_partial_failure(monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    monkeypatch.setattr(cli, "execute", lambda _args: {
        "success": False, "processed": 1, "failed": 1, "failed_rows": [2],
    })
    assert cli.main([]) == 1
    assert json.loads(capsys.readouterr().out)["failed_rows"] == [2]


def test_cli_rejects_arguments_before_workbook_access(capsys: pytest.CaptureFixture[str]) -> None:
    assert cli.main(["unexpected"]) == 1
    assert json.loads(capsys.readouterr().out)["error"] == "MMISClientError"


def test_cli_rejects_invalid_config_without_creating_a_reader(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> None:
    def invalid_config(cls: type) -> object:
        raise MMISClientError(".env 缺少 MMIS_USERNAME 或 MMIS_PASSWORD")

    monkeypatch.setattr(cli.MMISConfig, "from_env", classmethod(invalid_config))
    monkeypatch.setattr(
        cli, "PeriodicInspectionWorkOrderDetailReader",
        lambda _session: pytest.fail("Invalid config must stop before query"),
    )
    assert cli.main([]) == 1
    assert json.loads(capsys.readouterr().out)["error"] == "MMISClientError"


def test_cli_runs_batch_and_reports_the_real_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> None:
    work_order = "115-3A-100-001"
    source, output = make_files(tmp_path, [
        ("EMU801", "3A", work_order, "2026/10/20"),
    ])
    monkeypatch.setattr(cli, "DEFAULT_SOURCE", source)
    monkeypatch.setattr(cli, "DEFAULT_OUTPUT", output)
    monkeypatch.setattr(cli.MMISConfig, "from_env", classmethod(lambda cls: object()))
    monkeypatch.setattr(cli, "MMISSession", lambda _config: object())

    class Reader:
        def __init__(self, _session: object) -> None:
            pass

        def run_air_hose_dates(self, number: str) -> dict[str, object]:
            return query_result("EMU801", number, count=12)

    monkeypatch.setattr(cli, "PeriodicInspectionWorkOrderDetailReader", Reader)
    assert cli.main([]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["success"] is True
    assert result["measurements"] == 12
    assert result["output_file"] == str(output)
    assert len(visible_rows(output)) == 13
