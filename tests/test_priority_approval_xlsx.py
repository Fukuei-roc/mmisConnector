from __future__ import annotations

import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mmis_connector.auth import MMISClientError
from tools.mmis_development import _priority_approval_xlsx as priority
from tools.mmis_development import auto_process_unclosed_ab_fault_notices as snapshot


@pytest.fixture
def source(monkeypatch, tmp_path):
    database = tmp_path / "notices.db"
    columns = tuple(column[0] for column in snapshot.COLUMNS)
    with sqlite3.connect(database) as connection:
        connection.execute(snapshot.CREATE_TABLE_SQL)
    monkeypatch.setenv(priority.APPROVER_ENV_NAME, "新任組長")
    monkeypatch.setattr(priority, "load_dotenv", lambda *args, **kwargs: None)
    monkeypatch.setattr(priority.MMISConfig, "from_env", lambda: object())
    monkeypatch.setattr(priority, "MMISSession", lambda config: object())
    return database, columns, tmp_path


def insert_pending(database, *values):
    with sqlite3.connect(database) as connection:
        connection.executemany(
            'INSERT INTO unclosed_ab_fault_notices ("通報號", "待簽核工單") VALUES (?, ?)',
            [(str(index), json.dumps(value)) for index, value in enumerate(values)],
        )


def test_deduplicates_and_filters_both_work_order_types_and_matches_style(source, monkeypatch):
    database, columns, output = source
    insert_pending(database,
                   ["115-C1-1", "115-C2-2", "115-C3-3", "115-1A-4"],
                   ["115-C1-1", "115-1A-5", "115-1A-6"])
    calls = []

    def temporary(self, number):
        calls.append((number, self.client))
        return {"success": True, "工作單": number,
                "待簽核人": ["新任組長"] if number != "115-C2-2" else ["其他人"]}

    def daily(self, number):
        calls.append((number, self.client))
        return {"success": True, "工作單": number,
                "待簽核人": ["新任組長"] if number != "115-1A-6" else [],
                "檢修段": "臺北機務段" if number == "115-1A-5" else "新竹機務段"}

    monkeypatch.setattr(priority.TemporaryRepairProcedureReader, "run", temporary)
    monkeypatch.setattr(priority.DailyInspectionInspectionRecordReader, "run", daily)
    fixed = datetime(2026, 10, 31, tzinfo=timezone.utc)
    result = priority.export_priority_xlsx(database, snapshot.TABLE_NAME, columns,
                                           output_directory=output, now=fixed)
    assert len(calls) == 6 and len({id(client) for _, client in calls}) == 1
    assert result["records_exported"] == 3
    assert Path(result["output_file"]).name == "優先簽核工單_程式產生1031.xlsx"
    workbook = load_workbook(result["output_file"])
    try:
        sheet = workbook.active
        assert sheet.title == "優先簽核工單"
        assert list(sheet.values) == [("工作單號",), ("115-C1-1",), ("115-C3-3",), ("115-1A-4",)]
        assert sheet["A1"].font.bold and sheet["A1"].alignment.horizontal == "center"
        for row in sheet.iter_rows():
            cell = row[0]
            assert cell.font.name == "微軟正黑體" and cell.font.sz == 12
            assert cell.alignment.vertical == "top"
            assert sheet.row_dimensions[cell.row].height == 15.5
        assert sheet["A2"].number_format == "@" and sheet["A2"].data_type == "s"
        assert sheet["A2"].alignment.horizontal == "left"
    finally:
        workbook.close()


@pytest.mark.parametrize("number", ["115-C3-41264", "115-C12-41264", "115-C1-41264-001"])
def test_temporary_classes_and_test_run_children_use_temporary_reader(source, monkeypatch, number):
    database, columns, output = source
    insert_pending(database, [number])
    calls = []

    def temporary(self, value):
        calls.append(value)
        return {"success": True, "工作單": value, "待簽核人": ["新任組長"]}

    monkeypatch.setattr(priority.TemporaryRepairProcedureReader, "run", temporary)
    result = priority.export_priority_xlsx(database, snapshot.TABLE_NAME, columns,
                                          output_directory=output)
    assert calls == [number] and result["records_exported"] == 1


def test_empty_list_still_exports_header_and_does_not_overwrite(source):
    database, columns, output = source
    fixed = datetime(2026, 10, 31, tzinfo=timezone.utc)
    first = priority.export_priority_xlsx(database, snapshot.TABLE_NAME, columns,
                                          output_directory=output, now=fixed)
    second = priority.export_priority_xlsx(database, snapshot.TABLE_NAME, columns,
                                           output_directory=output, now=fixed)
    assert first["records_exported"] == second["records_exported"] == 0
    assert Path(second["output_file"]).name.endswith("1031_2.xlsx")
    workbook = load_workbook(first["output_file"])
    assert list(workbook.active.values) == [("工作單號",)]
    workbook.close()


@pytest.mark.parametrize("response", [
    {"success": False, "工作單": "115-C1-1", "待簽核人": []},
    {"success": True, "工作單": "115-C1-2", "待簽核人": []},
    {"success": True, "工作單": "115-C1-1", "待簽核人": "新任組長"},
])
def test_invalid_query_result_never_exports_partial_list(source, monkeypatch, response):
    database, columns, output = source
    insert_pending(database, ["115-C1-1"])
    monkeypatch.setattr(priority.TemporaryRepairProcedureReader, "run", lambda *args: response)
    with pytest.raises(MMISClientError):
        priority.export_priority_xlsx(database, snapshot.TABLE_NAME, columns, output_directory=output)
    assert list(output.glob("*.xlsx")) == []


def test_missing_approver_is_explicit_error(source, monkeypatch):
    database, columns, output = source
    monkeypatch.delenv(priority.APPROVER_ENV_NAME)
    with pytest.raises(MMISClientError, match=priority.APPROVER_ENV_NAME):
        priority.export_priority_xlsx(database, snapshot.TABLE_NAME, columns, output_directory=output)


@pytest.mark.parametrize("invalid", ["bad JSON", '"115-C1-1"', '[1]', '[""]'])
def test_invalid_database_pending_orders_are_rejected(invalid):
    with pytest.raises(MMISClientError):
        priority.pending_numbers([(invalid,)], ("待簽核工單",))
