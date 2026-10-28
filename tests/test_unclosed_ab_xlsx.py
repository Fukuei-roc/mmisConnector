from __future__ import annotations

import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mmis_connector.auth import MMISClientError
from mmis_connector.fault_notices.analysis_readiness import AnalysisEvidenceIncomplete
from mmis_connector.fault_notices.query import (
    UNCLOSED_DEPOT_NAME, UNCLOSED_INCIDENT_LEVELS, UNCLOSED_QUERY_NAME,
)
from tools.mmis_development import auto_process_unclosed_ab_fault_notices as tool
from tools.mmis_development import export_unclosed_ab_fault_notices_from_db as offline
from tools.mmis_development import _unclosed_ab_xlsx as xlsx


def source(number: str, *, phenomenon: str = "原始現象") -> dict:
    return {
        "車次": "123", "車組/車號": "EMU900", "發生日期": "2026/10/20",
        "事故等級": "B", "ATP故障": True, "故障現象": phenomenon,
        "狀態": "處理中", "通報號": number,
    }


def result(*records: dict) -> dict:
    return {
        "success": True, "query_name": UNCLOSED_QUERY_NAME,
        "filters": {"配屬段別名稱": UNCLOSED_DEPOT_NAME, "事故等級": UNCLOSED_INCIDENT_LEVELS},
        "count": len(records), "records": list(records),
    }


def detail(number: str, *, analysis_completed: bool = True) -> dict:
    return {
        "通報號": number,
        "結案狀態檢查": {
            "Sufficient Data for Fault Analysis": None,
            "Fault Analysis Completed": analysis_completed,
            "All Depot Maintenance Work Orders Closed": False,
            "Depot Maintenance Work Orders Pending Approval": False,
            "CA Repair Work Orders Not Closed": False,
            "ATP Fault Analysis Incomplete": False,
        },
    }


def setup_offline(monkeypatch, tmp_path: Path, *records: dict) -> tuple[Path, Path]:
    database = tmp_path / "data" / "unclosed_ab_fault_notices.db"
    output = tmp_path / "中文 Excel 資料夾"
    output.mkdir()
    monkeypatch.setattr(tool, "DATABASE_PATH", database)
    monkeypatch.setattr(tool, "fetch_unclosed_fault_notices", lambda: result(*records))
    monkeypatch.setattr(tool, "fetch_fault_notice_detail",
                        lambda number, **kwargs: detail(number))
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))
    return database, output


def columns() -> tuple[str, ...]:
    return tuple(column[0] for column in tool.COLUMNS)


def bool_columns() -> frozenset[str]:
    return frozenset({"ATP故障", *tool.DETAIL_CHECK_FIELDS})


def test_no_flag_syncs_without_xlsx(monkeypatch, tmp_path, capsys) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("new"))
    assert tool.main([]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["status"] == "success"
    assert "xlsx_export" not in summary
    assert database.is_file()
    assert list(output.iterdir()) == []
    with sqlite3.connect(database) as connection:
        assert connection.execute(
            'SELECT "備註" FROM unclosed_ab_fault_notices'
        ).fetchone()[0] == "段修工單未處理完成"


def test_export_syncs_fresh_sqlite_before_reading_xlsx(monkeypatch, tmp_path, capsys) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("new"))
    assert not database.exists()
    assert tool.main(["--export-xlsx"]) == 0
    summary = json.loads(capsys.readouterr().out)
    exported = Path(summary["xlsx_export"]["output_file"])
    assert summary["operation"] == "sync_and_export_xlsx"
    assert summary["records_saved"] == summary["xlsx_export"]["records_exported"] == 1
    assert summary["xlsx_export"]["columns_exported"] == 9
    assert exported.parent == output
    with sqlite3.connect(database) as connection:
        assert connection.execute('SELECT "通報號" FROM unclosed_ab_fault_notices').fetchone()[0] == "new"
    workbook = load_workbook(exported)
    assert workbook.sheetnames == [xlsx.SHEET_NAME]
    sheet = workbook.active
    assert tuple(cell.value for cell in sheet[1]) == tuple(
        name for name in columns() if name not in tool.DETAIL_CHECK_FIELDS
    )
    assert sheet.max_column == 9 and sheet.max_row == 2
    assert sheet.cell(2, 8).value == "new"
    assert sheet.cell(2, 8).fill.patternType is None
    assert sheet.cell(2, 5).value is True
    assert sheet.cell(2, 9).value == "段修工單未處理完成"
    assert sheet.row_dimensions[1].height == sheet.row_dimensions[2].height == 15.5
    for cell in sheet[1]:
        assert cell.font.name == xlsx.FONT_NAME and cell.font.sz == 12
        assert cell.font.bold and cell.alignment.horizontal == "center"
        assert cell.alignment.vertical == "top"
    for cell in sheet[2]:
        assert cell.font.name == xlsx.FONT_NAME and cell.font.sz == 12
        assert not cell.font.bold and cell.alignment.horizontal == "left"
        assert cell.alignment.vertical == "top"
    workbook.close()


def test_existing_sqlite_is_replaced_before_export(monkeypatch, tmp_path, capsys) -> None:
    database, _ = setup_offline(monkeypatch, tmp_path, source("new"))
    tool.replace_rows(tool.prepare_rows(result(source("old"))), database, tool.SyncState())
    assert tool.main(["--export-xlsx"]) == 0
    summary = json.loads(capsys.readouterr().out)
    workbook = load_workbook(summary["xlsx_export"]["output_file"])
    assert [row[7] for row in workbook.active.iter_rows(min_row=2, values_only=True)] == ["new"]
    workbook.close()


def test_export_replaces_unmarked_note_with_database_status_note(monkeypatch, tmp_path, capsys) -> None:
    database, _ = setup_offline(monkeypatch, tmp_path, source("one"))
    tool.replace_rows(tool.prepare_rows(result(source("one"))), database, tool.SyncState())
    with sqlite3.connect(database) as connection:
        connection.execute(
            'UPDATE unclosed_ab_fault_notices SET "備註" = ?',
            ("填寫故障分析完成",),
        )

    assert tool.main(["--export-xlsx"]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["fault_analysis_candidates"] == 0
    assert summary["xlsx_export"]["status"] == "success"
    with sqlite3.connect(database) as connection:
        assert connection.execute(
            'SELECT "備註" FROM unclosed_ab_fault_notices'
        ).fetchone()[0] == "段修工單未處理完成"


@pytest.mark.parametrize("value,expected", [(1, True), (0, False), (None, None)])
def test_boolean_and_null_from_sqlite(monkeypatch, tmp_path, value, expected) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("one"))
    tool.replace_rows(tool.prepare_rows(result(source("one"))), database, tool.SyncState())
    with sqlite3.connect(database) as connection:
        connection.execute(
            'UPDATE unclosed_ab_fault_notices SET "ATP故障"=?, '
            '"Fault Analysis Completed"=?', (value, value),
        )
    exported = xlsx.export_xlsx(database, tool.TABLE_NAME, columns(), bool_columns(),
                                output_directory=output)
    workbook = load_workbook(exported["output_file"])
    assert workbook.active.cell(2, 5).value is expected
    assert workbook.active.cell(2, 10).value is expected
    assert workbook.active.cell(2, 15).value is None
    workbook.close()


def test_text_formula_safety_width_and_wrapping(monkeypatch, tmp_path) -> None:
    text = "危險文字" * 70 + "\n第二行" * 30
    database, output = setup_offline(monkeypatch, tmp_path, source("one", phenomenon=text))
    tool.replace_rows(tool.prepare_rows(result(source("one", phenomenon=text))),
                      database, tool.SyncState())
    with sqlite3.connect(database) as connection:
        connection.execute(
            'UPDATE unclosed_ab_fault_notices SET "備註"=?, "車次"=?',
            ('=HYPERLINK("https://example.test","x")', "+SUM(1,2)"),
        )
    exported = xlsx.export_xlsx(database, tool.TABLE_NAME, columns(), bool_columns(),
                                output_directory=output)
    workbook = load_workbook(exported["output_file"])
    sheet = workbook.active
    assert sheet.cell(2, 1).value == "+SUM(1,2)"
    assert sheet.cell(2, 1).data_type == "s"
    assert sheet.cell(2, 15).value.startswith("=HYPERLINK")
    assert sheet.cell(2, 15).data_type == "s"
    assert sheet.cell(2, 6).value == text
    assert sheet.cell(2, 6).alignment.wrap_text is True
    assert sheet.column_dimensions["F"].width <= 65
    assert sheet.row_dimensions[2].height > 24
    assert sheet.column_dimensions["I"].width > sheet.column_dimensions["A"].width
    workbook.close()
    with sqlite3.connect(database) as connection:
        assert connection.execute('SELECT "車次" FROM unclosed_ab_fault_notices').fetchone()[0] == "+SUM(1,2)"


def test_date_suffix_and_no_overwrite_even_with_concurrent_exports(monkeypatch, tmp_path) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("one"))
    tool.replace_rows(tool.prepare_rows(result(source("one"))), database, tool.SyncState())
    fixed = datetime(2026, 11, 4, 16, 30, tzinfo=timezone.utc)
    first = output / "AB級未處理故障通報_程式產生1105.xlsx"
    first.write_bytes(b"existing")

    def export(_: int) -> dict:
        return xlsx.export_xlsx(database, tool.TABLE_NAME, columns(), bool_columns(),
                                output_directory=output, now=fixed)

    with ThreadPoolExecutor(max_workers=2) as pool:
        files = [Path(item["output_file"]) for item in pool.map(export, range(2))]
    assert {path.name for path in files} == {
        "AB級未處理故障通報_程式產生1105_2.xlsx",
        "AB級未處理故障通報_程式產生1105_3.xlsx",
    }
    assert first.read_bytes() == b"existing"
    for path in files:
        workbook = load_workbook(path)
        assert workbook.active.cell(2, 8).value == "one"
        workbook.close()


def test_first_filename_has_no_suffix_and_empty_snapshot_has_header(monkeypatch, tmp_path, capsys) -> None:
    _, output = setup_offline(monkeypatch, tmp_path)
    assert tool.main(["--export-xlsx"]) == 0
    summary = json.loads(capsys.readouterr().out)
    exported = Path(summary["xlsx_export"]["output_file"])
    assert exported.parent == output
    assert exported.name.startswith(xlsx.FILE_PREFIX)
    assert "_2.xlsx" not in exported.name
    workbook = load_workbook(exported)
    assert workbook.active.max_row == 1
    assert tuple(cell.value for cell in workbook.active[1]) == tuple(
        name for name in columns() if name not in tool.DETAIL_CHECK_FIELDS
    )
    workbook.close()


def test_sync_delegates_export_to_standalone_command(monkeypatch, tmp_path, capsys) -> None:
    database, _ = setup_offline(monkeypatch, tmp_path, source("one"))
    calls = []

    def export(args, *, database_path, expected_records):
        calls.append((args, database_path, expected_records))
        return {"xlsx_export": {"status": "success", "records_exported": 1}}

    monkeypatch.setattr(offline, "execute", export)
    assert tool.main(["--export-xlsx"]) == 0
    assert calls == [([], database, 1)]
    assert json.loads(capsys.readouterr().out)["xlsx_export"]["records_exported"] == 1


def test_quoted_env_path_and_missing_or_relative_setting(monkeypatch, tmp_path) -> None:
    output = tmp_path / "中文 with spaces"
    output.mkdir()
    monkeypatch.setattr(xlsx, "load_dotenv", lambda *args, **kwargs: None)
    monkeypatch.setenv(xlsx.ENV_NAME, f'"{output}"')
    assert xlsx.configured_output_directory() == output
    monkeypatch.setenv(xlsx.ENV_NAME, " ")
    with pytest.raises(xlsx.XlsxExportError, match="未設定"):
        xlsx.configured_output_directory()
    monkeypatch.setenv(xlsx.ENV_NAME, "relative/folder")
    with pytest.raises(xlsx.XlsxExportError, match="絕對路徑"):
        xlsx.configured_output_directory()


def test_query_failure_and_sqlite_failure_never_export(monkeypatch, tmp_path, capsys) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("new"))
    tool.replace_rows(tool.prepare_rows(result(source("old"))), database, tool.SyncState())
    monkeypatch.setattr(tool, "fetch_unclosed_fault_notices",
                        lambda: (_ for _ in ()).throw(MMISClientError("查詢失敗")))
    assert tool.main(["--export-xlsx"]) == 1
    assert json.loads(capsys.readouterr().out)["xlsx_export"]["status"] == "not_run"
    assert list(output.iterdir()) == []
    monkeypatch.setattr(tool, "fetch_unclosed_fault_notices", lambda: result(source("new")))
    with sqlite3.connect(database) as connection:
        connection.execute('''CREATE TRIGGER reject_new BEFORE INSERT ON unclosed_ab_fault_notices
                              BEGIN SELECT RAISE(FAIL, 'test'); END''')
    assert tool.main(["--export-xlsx"]) == 1
    assert json.loads(capsys.readouterr().out)["xlsx_export"]["status"] == "not_run"
    assert list(output.iterdir()) == []
    with sqlite3.connect(database) as connection:
        assert connection.execute('SELECT "通報號" FROM unclosed_ab_fault_notices').fetchone()[0] == "old"


def test_sqlite_commit_failure_never_exports_old_snapshot(monkeypatch, tmp_path, capsys) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("new"))
    tool.replace_rows(tool.prepare_rows(result(source("old"))), database, tool.SyncState())
    real_connect = sqlite3.connect

    class FailCommit:
        def __init__(self, connection):
            self.connection = connection

        def __getattr__(self, name):
            return getattr(self.connection, name)

        def commit(self):
            raise sqlite3.OperationalError("commit rejected")

    def connect(path, *args, **kwargs):
        connection = real_connect(path, *args, **kwargs)
        return FailCommit(connection) if Path(path) == database else connection

    with monkeypatch.context() as patch:
        patch.setattr(tool.sqlite3, "connect", connect)
        assert tool.main(["--export-xlsx"]) == 1
    summary = json.loads(capsys.readouterr().out)
    assert summary["rolled_back"] is True
    assert summary["xlsx_export"]["status"] == "not_run"
    assert list(output.iterdir()) == []
    with sqlite3.connect(database) as connection:
        assert connection.execute('SELECT "通報號" FROM unclosed_ab_fault_notices').fetchone()[0] == "old"


def test_export_failure_keeps_committed_sqlite(monkeypatch, tmp_path, capsys) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("new"))
    tool.replace_rows(tool.prepare_rows(result(source("old"))), database, tool.SyncState())
    monkeypatch.setenv(xlsx.ENV_NAME, str(output / "missing"))
    assert tool.main(["--export-xlsx"]) == 1
    summary = json.loads(capsys.readouterr().out)
    assert summary["sqlite_sync_success"] is True
    assert summary["xlsx_export"]["status"] == "failed"
    with sqlite3.connect(database) as connection:
        assert connection.execute('SELECT "通報號" FROM unclosed_ab_fault_notices').fetchone()[0] == "new"
    assert list(output.iterdir()) == []


def test_missing_export_setting_fails_only_after_sync(monkeypatch, tmp_path, capsys) -> None:
    database, _ = setup_offline(monkeypatch, tmp_path, source("new"))
    monkeypatch.delenv(xlsx.ENV_NAME)
    monkeypatch.setattr(xlsx, "load_dotenv", lambda *args, **kwargs: None)
    assert tool.main(["--export-xlsx"]) == 1
    summary = json.loads(capsys.readouterr().out)
    assert summary["sqlite_sync_success"] is True
    assert xlsx.ENV_NAME in summary["message"]
    assert database.is_file()
    assert tool.main([]) == 0


def test_partial_sync_skips_export(monkeypatch, tmp_path, capsys) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("one"))

    def partial(number, *, check_analysis_readiness):
        if check_analysis_readiness:
            raise AnalysisEvidenceIncomplete("工單證據不完整")
        return detail(number, analysis_completed=False)

    monkeypatch.setattr(tool, "fetch_fault_notice_detail", partial)
    assert tool.main(["--export-xlsx"]) == 1
    summary = json.loads(capsys.readouterr().out)
    assert summary["status"] == "partial_success"
    assert summary["xlsx_export"]["status"] == "skipped"
    assert database.is_file() and list(output.iterdir()) == []


def test_failed_workbook_save_cleans_only_new_file(monkeypatch, tmp_path) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("one"))
    tool.replace_rows(tool.prepare_rows(result(source("one"))), database, tool.SyncState())
    existing = output / "keep.xlsx"
    existing.write_bytes(b"existing")

    def fail_save(self, handle):
        handle.write(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(xlsx.Workbook, "save", fail_save)
    with pytest.raises(xlsx.XlsxExportError, match="XLSX 寫入失敗"):
        xlsx.export_xlsx(database, tool.TABLE_NAME, columns(), bool_columns(),
                         output_directory=output)
    assert list(output.iterdir()) == [existing]
    assert existing.read_bytes() == b"existing"


def test_xlsx_requires_matching_schema_and_valid_boolean(monkeypatch, tmp_path) -> None:
    database, output = setup_offline(monkeypatch, tmp_path, source("one"))
    tool.replace_rows(tool.prepare_rows(result(source("one"))), database, tool.SyncState())
    with pytest.raises(xlsx.XlsxExportError, match="欄位"):
        xlsx.export_xlsx(database, tool.TABLE_NAME, ("wrong",), bool_columns(),
                         output_directory=output)
    with sqlite3.connect(database) as connection:
        connection.execute('UPDATE unclosed_ab_fault_notices SET "ATP故障"=2')
    with pytest.raises(xlsx.XlsxExportError, match="布林"):
        xlsx.export_xlsx(database, tool.TABLE_NAME, columns(), bool_columns(),
                         output_directory=output)
    assert list(output.iterdir()) == []
