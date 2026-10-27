from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mmis_connector.auth import MMISClientError
from mmis_connector.periodic_inspection import workbook as batch
from tools.mmis_development import fill_latest_periodic_inspection_work_orders as cli


def make_workbook(path: Path, vehicles: list[str], headers: tuple[str, ...] = batch.HEADERS) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = batch.SHEET_NAME
    sheet.append(headers)
    for vehicle in vehicles:
        sheet.append([vehicle])
    workbook.save(path)


def result(vehicle: str, work_order: str = "115-2A-1") -> dict[str, str]:
    return {
        "車組/車號": vehicle,
        "檢修級別": "2A",
        "工作單": work_order,
        "檢修日期": "2026/10/27",
    }


def read_rows(path: Path) -> list[tuple[object, ...]]:
    workbook = load_workbook(path)
    try:
        return list(workbook[batch.SHEET_NAME].values)
    finally:
        workbook.close()


def test_validates_file_and_headers_before_query(tmp_path: Path) -> None:
    with pytest.raises(MMISClientError, match=".xlsx"):
        batch.load_validated_workbook(tmp_path / "missing.xlsx")
    path = tmp_path / "wrong.xlsx"
    make_workbook(path, ["813"], ("車組/車號", "錯誤", "工作單", "檢修日期"))
    with pytest.raises(MMISClientError, match="標題"):
        batch.load_validated_workbook(path)
    assert read_rows(path)[1] == ("813", None, None, None)


def test_retries_failed_row_with_fresh_query_and_saves_success(tmp_path: Path) -> None:
    path = tmp_path / "vehicles.xlsx"
    make_workbook(path, ["813", "814"])
    queries_created = 0
    calls: list[str] = []

    class Query:
        def run(self, vehicle: str) -> dict[str, str]:
            calls.append(vehicle)
            if vehicle == "813" and len(calls) == 1:
                raise MMISClientError("MMIS 網路請求失敗")
            return result(vehicle)

    def factory() -> Query:
        nonlocal queries_created
        queries_created += 1
        return Query()

    workbook = batch.load_validated_workbook(path)
    summary = batch.fill_latest_periodic_inspection_work_orders(workbook, path, factory)
    workbook.close()

    assert summary.as_dict() == {
        "success": True, "processed": 2, "skipped": 0, "failed": 0, "failed_rows": []
    }
    assert calls == ["813", "813", "814"]
    assert queries_created == 2
    assert read_rows(path)[1:] == [
        ("813", "2A", "115-2A-1", "2026/10/27"),
        ("814", "2A", "115-2A-1", "2026/10/27"),
    ]


def test_exhausted_retries_continue_and_rerun_only_incomplete_rows(tmp_path: Path) -> None:
    path = tmp_path / "vehicles.xlsx"
    make_workbook(path, ["813", "814", "815"])
    calls: list[str] = []

    class FirstQuery:
        def run(self, vehicle: str) -> dict[str, str]:
            calls.append(vehicle)
            if vehicle == "814":
                raise MMISClientError("網路中斷")
            return result(vehicle)

    workbook = batch.load_validated_workbook(path)
    first = batch.fill_latest_periodic_inspection_work_orders(workbook, path, FirstQuery)
    workbook.close()
    assert first.as_dict() == {
        "success": False, "processed": 2, "skipped": 0, "failed": 1, "failed_rows": [3]
    }
    assert calls == ["813", "814", "814", "814", "815"]
    assert read_rows(path)[2] == ("814", None, None, None)

    class SecondQuery:
        def run(self, vehicle: str) -> dict[str, str]:
            calls.append(vehicle)
            return result(vehicle)

    workbook = batch.load_validated_workbook(path)
    second = batch.fill_latest_periodic_inspection_work_orders(workbook, path, SecondQuery)
    workbook.close()
    assert second.as_dict() == {
        "success": True, "processed": 1, "skipped": 2, "failed": 0, "failed_rows": []
    }
    assert calls[-1] == "814"
    assert read_rows(path)[2] == ("814", "2A", "115-2A-1", "2026/10/27")


def test_blank_vehicle_and_complete_row_are_skipped(tmp_path: Path) -> None:
    path = tmp_path / "vehicles.xlsx"
    make_workbook(path, ["", "813"])
    workbook = load_workbook(path)
    sheet = workbook[batch.SHEET_NAME]
    for column, value in enumerate(("2A", "old", "2026/01/01"), start=2):
        sheet.cell(3, column).value = value
    workbook.save(path)
    workbook.close()
    workbook = batch.load_validated_workbook(path)
    summary = batch.fill_latest_periodic_inspection_work_orders(
        workbook, path, lambda: pytest.fail("query must not run")
    )
    workbook.close()
    assert summary.skipped == 2
    assert read_rows(path)[2] == ("813", "2A", "old", "2026/01/01")


def test_atomic_replace_failure_keeps_original_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "vehicles.xlsx"
    make_workbook(path, ["813"])

    class Query:
        def run(self, vehicle: str) -> dict[str, str]:
            return result(vehicle)

    def fail_replace(source: Path, destination: Path) -> None:
        raise PermissionError("locked")

    monkeypatch.setattr(batch.os, "replace", fail_replace)
    workbook = batch.load_validated_workbook(path)
    with pytest.raises(MMISClientError, match="無法儲存"):
        batch.fill_latest_periodic_inspection_work_orders(workbook, path, Query)
    workbook.close()
    assert read_rows(path)[1] == ("813", None, None, None)
    assert list(tmp_path.glob(".*.xlsx")) == []


def test_work_order_is_text_even_when_it_starts_with_formula_marker(tmp_path: Path) -> None:
    path = tmp_path / "vehicles.xlsx"
    make_workbook(path, ["813"])

    class Query:
        def run(self, vehicle: str) -> dict[str, str]:
            return result(vehicle, "=FORMULA()")

    workbook = batch.load_validated_workbook(path)
    batch.fill_latest_periodic_inspection_work_orders(workbook, path, Query)
    workbook.close()
    saved = load_workbook(path)
    try:
        cell = saved[batch.SHEET_NAME]["C2"]
        assert cell.value == "=FORMULA()"
        assert cell.data_type == "s"
    finally:
        saved.close()


def test_cli_reports_partial_failure_with_nonzero_exit(monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    monkeypatch.setattr(cli, "execute", lambda args: batch.BatchResult(2, 0, (4,)).as_dict())
    assert cli.main([]) == 1
    assert '"failed_rows": [\n    4\n  ]' in capsys.readouterr().out
