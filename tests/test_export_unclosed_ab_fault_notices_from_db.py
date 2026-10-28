from __future__ import annotations

import builtins
import hashlib
import json
import sqlite3

import pytest

from openpyxl import load_workbook

from tools.mmis_development import auto_process_unclosed_ab_fault_notices as snapshot
from tools.mmis_development import export_unclosed_ab_fault_notices_from_db as offline
from tools.mmis_development import _unclosed_ab_xlsx as xlsx


def test_expected_records_rejects_changed_snapshot(monkeypatch, tmp_path) -> None:
    database = tmp_path / "unclosed_ab_fault_notices.db"
    output = tmp_path / "xlsx"
    output.mkdir()
    with sqlite3.connect(database) as connection:
        connection.execute(snapshot.CREATE_TABLE_SQL)
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    with pytest.raises(xlsx.XlsxExportError, match="筆數"):
        offline.execute([], database_path=database, expected_records=1)
    assert list(output.iterdir()) == []


def test_exports_existing_database_without_sync(monkeypatch, tmp_path, capsys) -> None:
    database = tmp_path / "unclosed_ab_fault_notices.db"
    output = tmp_path / "xlsx"
    output.mkdir()
    with sqlite3.connect(database) as connection:
        connection.execute(snapshot.CREATE_TABLE_SQL)
        connection.execute(
            snapshot.INSERT_SQL,
            ("123", "EMU900", "2026/10/21", "B", 1, "測試故障", "處理中", "1151021-01"),
        )
        connection.execute(
            'UPDATE unclosed_ab_fault_notices SET "Fault Analysis Completed" = 0 '
            'WHERE "通報號" = ?',
            ("1151021-01",),
        )
    before = hashlib.sha256(database.read_bytes()).digest()
    monkeypatch.setattr(snapshot, "DATABASE_PATH", database)
    monkeypatch.setattr(
        snapshot, "fetch_unclosed_fault_notices",
        lambda: (_ for _ in ()).throw(AssertionError("MMIS query must not run")),
    )
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    assert offline.main([]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["operation"] == "export_existing_sqlite_to_xlsx"
    assert result["xlsx_export"]["records_exported"] == 1
    assert result["xlsx_export"]["columns_exported"] == 9
    assert hashlib.sha256(database.read_bytes()).digest() == before

    workbook = load_workbook(result["xlsx_export"]["output_file"])
    try:
        sheet = workbook.active
        assert sheet.title == xlsx.SHEET_NAME
        assert tuple(cell.value for cell in sheet[1]) == (
            "車次", "車組/車號", "發生日期", "事故等級", "ATP故障",
            "故障現象", "狀態", "通報號", "備註",
        )
        assert sheet.cell(2, 8).value == "1151021-01"
        assert sheet.cell(2, 5).value is True
        assert sheet.cell(2, 9).value is None
        assert all(sheet.row_dimensions[row].height == 15.5 for row in (1, 2))
        assert all(
            cell.alignment.vertical == "top"
            for row in sheet.iter_rows()
            for cell in row
        )
        assert [sheet.column_dimensions[column].width for column in "ABCDEFGHI"] == [
            6.5, 11.8, 14.5, 10.6, 10.6, 70, 10.6, 14.9, 40,
        ]
    finally:
        workbook.close()


def test_sorts_all_rows_by_level_date_and_vehicle(monkeypatch, tmp_path, capsys) -> None:
    database = tmp_path / "unclosed_ab_fault_notices.db"
    output = tmp_path / "xlsx"
    output.mkdir()
    source_rows = [
        ("B", "2026/10/01", "EMU100", "B1"),
        ("A", "2026/10/02", "EMU300", "A3"),
        ("A", "2026/10/01", "EMU200", "A2"),
        ("A", "2026/10/01", "EMU100", "A1"),
        ("B", "2026/09/30", "EMU900", "B0"),
        ("A", "2026/10/01", "EMU100", "A1b"),
        ("A", None, "EMU400", "A-null"),
    ]
    with sqlite3.connect(database) as connection:
        connection.execute(snapshot.CREATE_TABLE_SQL)
        connection.executemany(
            snapshot.INSERT_SQL,
            [
                ("123", vehicle, date, level, 0, "故障", "處理中", number)
                for level, date, vehicle, number in source_rows
            ],
        )
    before = hashlib.sha256(database.read_bytes()).digest()
    monkeypatch.setattr(snapshot, "DATABASE_PATH", database)
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    assert offline.main([]) == 0
    result = json.loads(capsys.readouterr().out)
    workbook = load_workbook(result["xlsx_export"]["output_file"])
    try:
        assert [row[7] for row in workbook.active.iter_rows(min_row=2, values_only=True)] == [
            "A1", "A1b", "A2", "A3", "A-null", "B0", "B1",
        ]
    finally:
        workbook.close()
    assert hashlib.sha256(database.read_bytes()).digest() == before


def test_status_checks_write_matching_export_notes_only(monkeypatch, tmp_path, capsys) -> None:
    database = tmp_path / "unclosed_ab_fault_notices.db"
    output = tmp_path / "xlsx"
    output.mkdir()
    cases = [
        ("depot-open", (None, 1, 0, 0, 0, 0), "原備註", "段修工單未處理完成"),
        ("depot-approval", (None, 1, 0, 1, 0, 0), None, "等候段修工單簽核中"),
        ("pending-unprocessed", (None, 0, 0, 1, 0, 0), "原備註", "未處理"),
        ("checked-pending-unprocessed", (0, 0, 0, 1, 0, 0), "原備註", "未處理"),
        ("depot-closed-analysis-incomplete", (0, 0, 1, 0, 0, 0), "原備註", "未處理（段修工單結案）"),
        ("not-pending-unprocessed", (None, 0, 0, 0, 0, 0), "原備註", "原備註"),
        ("closable", (None, 1, 1, 0, 0, 0), None, "可結案"),
        ("unprocessed", (0, 0, 0, 0, 0, 0), None, "未處理"),
        ("different-sufficient", (1, 1, 0, 0, 0, 0), "=SUM(1,2)", "=SUM(1,2)"),
        ("different-ca", (None, 1, 0, 0, 1, 0), None, None),
        ("different-atp", (None, 1, 0, 0, 0, 1), None, None),
        ("unchecked", (None, None, None, None, None, None), "原備註", "原備註"),
    ]
    with sqlite3.connect(database) as connection:
        connection.execute(snapshot.CREATE_TABLE_SQL)
        for index, (number, status, note, _) in enumerate(cases):
            connection.execute(
                snapshot.INSERT_SQL,
                ("123", f"EMU{index}", f"2026/10/{21 - index:02d}",
                 "B", 0, "故障", "處理中", number),
            )
            connection.execute(
                'UPDATE unclosed_ab_fault_notices SET '
                '"Sufficient Data for Fault Analysis" = ?, '
                '"Fault Analysis Completed" = ?, '
                '"All Depot Maintenance Work Orders Closed" = ?, '
                '"Depot Maintenance Work Orders Pending Approval" = ?, '
                '"CA Repair Work Orders Not Closed" = ?, '
                '"ATP Fault Analysis Incomplete" = ?, "備註" = ? '
                'WHERE "通報號" = ?',
                (*status, note, number),
            )
    before = hashlib.sha256(database.read_bytes()).digest()
    monkeypatch.setattr(snapshot, "DATABASE_PATH", database)
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    assert offline.main([]) == 0
    result = json.loads(capsys.readouterr().out)
    workbook = load_workbook(result["xlsx_export"]["output_file"])
    try:
        sheet = workbook.active
        exported = {
            sheet.cell(row_number, 8).value: (
                sheet.cell(row_number, 8), sheet.cell(row_number, 9)
            )
            for row_number in range(2, sheet.max_row + 1)
        }
        assert all(cell.fill.patternType is None for row in sheet.iter_rows() for cell in row)
        for number, _, _, expected_note in cases:
            assert exported[number][1].value == expected_note
        assert exported["different-sufficient"][1].data_type == "s"
    finally:
        workbook.close()
    assert hashlib.sha256(database.read_bytes()).digest() == before


def test_missing_database_does_not_create_it(monkeypatch, tmp_path, capsys) -> None:
    database = tmp_path / "missing.db"
    output = tmp_path / "xlsx"
    output.mkdir()
    monkeypatch.setattr(snapshot, "DATABASE_PATH", database)
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    assert offline.main([]) == 1
    result = json.loads(capsys.readouterr().out)
    assert result["success"] is False
    assert not database.exists()
    assert list(output.iterdir()) == []


def test_invalid_hidden_status_rejects_export(monkeypatch, tmp_path, capsys) -> None:
    database = tmp_path / "invalid-status.db"
    output = tmp_path / "xlsx"
    output.mkdir()
    with sqlite3.connect(database) as connection:
        connection.execute(snapshot.CREATE_TABLE_SQL)
        connection.execute(
            snapshot.INSERT_SQL,
            ("123", "EMU900", "2026/10/21", "B", 0, "故障", "處理中", "1151021-01"),
        )
        connection.execute(
            'UPDATE unclosed_ab_fault_notices SET "ATP Fault Analysis Incomplete" = 2'
        )
    monkeypatch.setattr(snapshot, "DATABASE_PATH", database)
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    assert offline.main([]) == 1
    result = json.loads(capsys.readouterr().out)
    assert result["success"] is False
    assert "ATP Fault Analysis Incomplete" in result["message"]
    assert list(output.iterdir()) == []


def test_rejects_arguments_before_export(monkeypatch, tmp_path, capsys) -> None:
    output = tmp_path / "xlsx"
    output.mkdir()
    monkeypatch.setenv(xlsx.ENV_NAME, str(output))

    assert offline.main(["--export-xlsx"]) == 1
    result = json.loads(capsys.readouterr().out)
    assert "用法" in result["message"]
    assert list(output.iterdir()) == []


def test_missing_openpyxl_returns_json_error(monkeypatch, capsys) -> None:
    original_import = builtins.__import__

    def without_openpyxl(name, *args, **kwargs):
        if name == "_unclosed_ab_xlsx":
            raise ModuleNotFoundError("No module named 'openpyxl'", name="openpyxl")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_openpyxl)

    assert offline.main([]) == 1
    result = json.loads(capsys.readouterr().out)
    assert result["success"] is False
    assert result["error"] == "MMISClientError"
    assert "openpyxl" in result["message"]
